"""
convert_movies.py
讀取 export.xlsx 的電影工作表，輸出 public/data/movies.json

用法：
  python scripts/convert_movies.py                     # 使用預設路徑
  python scripts/convert_movies.py path/to/excel.xlsx  # 指定 Excel 路徑

讀取的工作表：
  每天電影排名資料        — 日榜（2021-04 起）
  Netflix 每週電影排名    — 週榜（2024-01 起）

設計規格：docs/superpowers/specs/2026-08-14-movie-page-design.md
"""

import json
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from board_common import build_coverage, find_sheet, majority_vote, to_iso_date
from language_rules import categorize

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_EXCEL = ROOT / "爬蟲臉書" / "output" / "export.xlsx"
OUT_PATH = ROOT / "public" / "data" / "movies.json"

SHEET_DAILY = "每天電影排名資料"
SHEET_WEEKLY = "Netflix 每週電影排名"

PENDING: Counter = Counter()
CONFLICTS: dict[str, set] = defaultdict(set)


def resolve_score(score, rank):
    """積分 = 11 - 排名。

    Excel 的積分欄是表格公式，openpyxl 以 data_only=True 讀取時公式快取
    可能已被清掉而讀成 None 或 0。與 convert_excel.py 的處理一致。
    """
    if score:
        return float(score)
    if rank is not None and 1 <= rank <= 10:
        return float(11 - rank)
    return 0.0


def header_index(ws) -> dict:
    """回傳 欄名 → 欄索引（0 起算）"""
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        return {str(v).strip(): i for i, v in enumerate(row) if v is not None}
    return {}


def parse_daily(ws):
    """回傳 (dailyBoard, raw_rows)

    dailyBoard: [{ date, entries: [{ rank, title }] }]，依日期升冪
    raw_rows:   [(title, raw_genre, is_original, iso_date, rank)]，供屬性彙整
    """
    cols = header_index(ws)
    by_date = defaultdict(list)
    raw_rows = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        iso = to_iso_date(row[cols["日期"]])
        title = str(row[cols["節目名稱"]] or "").strip()
        rank = row[cols["排名"]]
        if not iso or not title or rank is None:
            continue
        rank = int(rank)
        if not 1 <= rank <= 10:
            continue

        by_date[iso].append({"rank": rank, "title": title})
        raw_rows.append((
            title,
            row[cols["類型"]],
            bool(row[cols["是否Netflix Original"]]),
            iso,
            rank,
        ))

    daily_board = [
        {"date": iso, "entries": sorted(entries, key=lambda e: e["rank"])}
        for iso, entries in sorted(by_date.items())
    ]
    return daily_board, raw_rows


def parse_weekly(ws):
    """回傳 [{ weekNumber, dateRange, rankings: [{ rank, title, score }] }]"""
    cols = header_index(ws)
    by_week = defaultdict(list)

    for row in ws.iter_rows(min_row=2, values_only=True):
        start = to_iso_date(row[cols["日期起"]])
        end = to_iso_date(row[cols["日期迄"]])
        title = str(row[cols["節目名稱"]] or "").strip()
        rank = row[cols["排名"]]
        if not start or not title or rank is None:
            continue
        rank = int(rank)
        by_week[(start, end)].append({
            "rank": rank,
            "title": title,
            "score": resolve_score(row[cols["積分"]], rank),
        })

    weeks = []
    for i, ((start, end), items) in enumerate(sorted(by_week.items()), start=1):
        weeks.append({
            "weekNumber": i,
            "dateRange": f"{start} ~ {end}",
            "rankings": sorted(items, key=lambda e: e["rank"]),
        })
    return weeks


def build_entities(raw_rows) -> dict:
    """以多數決決定每部片的語言與形式，並彙整全期統計"""
    per_title = defaultdict(lambda: {
        "raw": [], "original": [], "dates": [], "ranks": [],
    })

    for title, raw_genre, is_original, iso, rank in raw_rows:
        bucket = per_title[title]
        bucket["raw"].append("" if raw_genre is None else str(raw_genre).strip())
        bucket["original"].append(is_original)
        bucket["dates"].append(iso)
        bucket["ranks"].append(rank)

    entities = {}
    for title, bucket in per_title.items():
        distinct = {r for r in bucket["raw"] if r}
        if len(distinct) > 1:
            CONFLICTS[title] = distinct

        winner = majority_vote(bucket["raw"])
        language, fmt, origin, reason = categorize(winner)
        if reason == "pending":
            PENDING[winner or "(空白)"] += len(bucket["raw"])

        ranks = bucket["ranks"]
        dates = sorted(bucket["dates"])
        entities[title] = {
            "language": language,
            "format": fmt,
            "origin": origin,
            "isNetflixOriginal": any(bucket["original"]),
            "firstDate": dates[0],
            "lastDate": dates[-1],
            "daysOnChart": len(dates),
            "bestRank": min(ranks),
            "avgRank": round(sum(ranks) / len(ranks), 2),
            "totalScore": sum(11 - r for r in ranks),
        }
    return entities


def print_report(entities, daily_board):
    total_slots = sum(len(b["entries"]) for b in daily_board)
    per_language = Counter()
    for board in daily_board:
        for entry in board["entries"]:
            attrs = entities.get(entry["title"])
            if attrs:
                per_language[attrs["language"]] += 1

    print("\n語言分佈（榜位天數）")
    for language, count in per_language.most_common():
        print(f"  {language:<6} {count:>6}  {count / total_slots * 100:5.1f}%")

    if PENDING:
        print(f"\n待審類型寫法（{len(PENDING)} 種，全部歸入「其他語言」）")
        for raw, count in PENDING.most_common(30):
            print(f"  {raw!r:<30} {count} 列")
        if len(PENDING) > 30:
            print(f"  …另有 {len(PENDING) - 30} 種")

    if CONFLICTS:
        print(f"\n類型前後不一致的片名（{len(CONFLICTS)} 部，已以多數決處理）")
        for title, variants in list(CONFLICTS.items())[:20]:
            print(f"  {title}: {sorted(variants)}")
        if len(CONFLICTS) > 20:
            print(f"  …另有 {len(CONFLICTS) - 20} 部")


def main():
    excel_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_EXCEL
    print(f"讀取 Excel：{excel_path}")
    if not excel_path.exists():
        print(f"ERROR: 找不到 {excel_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(excel_path, data_only=True)

    ws_daily = find_sheet(wb, SHEET_DAILY)
    if not ws_daily:
        print(f"ERROR: 找不到 sheet「{SHEET_DAILY}」")
        sys.exit(1)
    print(f"\n解析日榜：{SHEET_DAILY}")
    daily_board, raw_rows = parse_daily(ws_daily)
    print(f"  → {len(daily_board)} 天，{len(raw_rows)} 筆")

    ws_weekly = find_sheet(wb, SHEET_WEEKLY)
    weekly = []
    if ws_weekly:
        print(f"\n解析週榜：{SHEET_WEEKLY}")
        weekly = parse_weekly(ws_weekly)
        print(f"  → {len(weekly)} 週")
    else:
        print(f"  ⚠ 找不到 sheet「{SHEET_WEEKLY}」，週榜留空")

    print("\n彙整片名屬性")
    entities = build_entities(raw_rows)
    print(f"  → {len(entities)} 部電影")

    coverage = build_coverage([b["date"] for b in daily_board])
    print_report(entities, daily_board)

    output = {
        "meta": {
            "generatedAt": datetime.now().isoformat(),
            "dataThrough": daily_board[-1]["date"] if daily_board else "",
            "coverage": coverage,
        },
        "entities": entities,
        "dailyBoard": daily_board,
        "weeklyRankings": weekly,
    }

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, separators=(",", ":"))
    size_kb = OUT_PATH.stat().st_size // 1024
    print(f"\n輸出：{OUT_PATH} ({size_kb} KB)")

    print("\n逐年覆蓋率")
    for row in coverage:
        total = row["haveDays"] + row["missingDays"]
        print(f"  {row['year']}  {row['haveDays']:>3}/{total:<3} 缺 {row['missingDays']}")


if __name__ == "__main__":
    main()
